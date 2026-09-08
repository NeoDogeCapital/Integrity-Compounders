"""
export_universe_xlsx.py — full strategy universe workbook for review

Tabs: Universe (all active core names), Q1..Q4, Portfolio (holdings incl.
P&L vs avg cost). One row per name with the current model read: quads/axes,
QGS + tier, quality indicators + profile, earnings quality, alignment v3,
trend/extension/12-1 momentum, pillar scores, Fiscal-sourced fundamentals,
surprises and performance. House style header (navy, Calibri).

    python scripts/export_universe_xlsx.py [--out PATH]
"""
import sys
import argparse
from datetime import date
from pathlib import Path

ROOT = Path(__file__).parent.parent
sys.path.insert(0, str(ROOT))

import psycopg2
from config.settings import settings
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

COLS = [
    ("Ticker", "ticker", 9), ("Company", "company_name", 30), ("Sector", "sector", 20),
    ("Quad", "quadrant", 6), ("X RevMom", "x_rev_mom", 9), ("Y EpsMom", "x_eps_mom", 9),
    ("QGS Tier", "qgs_tier", 16), ("QGS", "quality_growth_score", 9),
    ("Quality Profile", "quality_profile", 16), ("Ind /6", "indicators_pass", 7),
    ("Earnings Quality", "earnings_quality_flag", 15),
    ("Align v3", "alignment_score_v3", 9), ("Bucket", "alignment_bucket_v3", 11),
    ("Trend", "trend_status", 10), ("Extension", "extension_flag", 12),
    ("Mom 12-1", "mom_12_1_risk_adj", 9),
    ("P1", "p1", 6), ("P2", "p2", 6), ("P3", "p3", 6), ("Composite", "comp", 9),
    ("ROIC %", "roic_trailing", 8), ("Gross Mgn %", "gross_margin_trailing", 10),
    ("Op Mgn %", "op_margin", 9), ("FCF Mgn %", "fcf_margin_trailing", 9),
    ("ND/EBITDA", "net_debt_ebitda", 9), ("FCF/EV", "fcf_ev_yield", 8),
    ("Fwd Rev %", "fwd_revenue_3y_cagr", 9), ("Fwd EPS %", "fwd_eps_3y_cagr", 9),
    ("Rev 3Y %", "revenue_3y_cagr_trailing", 9), ("EPS 3Y %", "eps_3y_cagr_trailing", 9),
    ("Buyback %", "buyback_yield", 9), ("EPS Surp %", "eps_surprise_q", 9),
    ("Rev Surp %", "rev_surprise_q", 9), ("Beta", "beta", 6),
    ("Mkt Cap $B", "mcap_b", 10), ("Price", "current_price", 9), ("YTD %", "ytd_perf", 8),
]
NAVY, GOLD = "1F3A5F", "C9A84C"


def fetch(cur):
    cur.execute("""SELECT c.ticker, c.company_name, c.sector, cmd.quadrant, cmd.x_rev_mom,
        cmd.x_eps_mom, cmd.qgs_tier, cmd.quality_growth_score, cmd.quality_profile,
        cmd.indicators_pass, cmd.earnings_quality_flag, cmd.alignment_score_v3,
        cmd.alignment_bucket_v3, cmd.trend_status, cmd.extension_flag, cmd.mom_12_1_risk_adj,
        cs.p1_business_quality, cs.p2_management, cs.p3_financial_strength, cs.composite_score_v2,
        cmd.roic_trailing, cmd.gross_margin_trailing, cmd.op_margin, cmd.fcf_margin_trailing,
        cmd.net_debt_ebitda, cmd.fcf_ev_yield, cmd.fwd_revenue_3y_cagr, cmd.fwd_eps_3y_cagr,
        cmd.revenue_3y_cagr_trailing, cmd.eps_3y_cagr_trailing, cmd.buyback_yield,
        cmd.eps_surprise_q, cmd.rev_surprise_q, cmd.beta, cmd.market_cap, cmd.current_price,
        cmd.ytd_perf, COALESCE(c.in_portfolio, FALSE)
        FROM companies c
        JOIN company_market_data cmd ON cmd.ticker = c.ticker
         AND cmd.data_date = (SELECT MAX(data_date) FROM company_market_data WHERE ticker = c.ticker)
        LEFT JOIN LATERAL (SELECT p1_business_quality, p2_management, p3_financial_strength,
            composite_score_v2 FROM company_scores WHERE company_id = c.id
            ORDER BY score_date DESC LIMIT 1) cs ON TRUE
        WHERE c.active = TRUE ORDER BY c.ticker""")
    names = ["ticker", "company_name", "sector", "quadrant", "x_rev_mom", "x_eps_mom",
             "qgs_tier", "quality_growth_score", "quality_profile", "indicators_pass",
             "earnings_quality_flag", "alignment_score_v3", "alignment_bucket_v3",
             "trend_status", "extension_flag", "mom_12_1_risk_adj", "p1", "p2", "p3", "comp",
             "roic_trailing", "gross_margin_trailing", "op_margin", "fcf_margin_trailing",
             "net_debt_ebitda", "fcf_ev_yield", "fwd_revenue_3y_cagr", "fwd_eps_3y_cagr",
             "revenue_3y_cagr_trailing", "eps_3y_cagr_trailing", "buyback_yield",
             "eps_surprise_q", "rev_surprise_q", "beta", "market_cap", "current_price",
             "ytd_perf", "in_portfolio"]
    rows = []
    for r in cur.fetchall():
        d = dict(zip(names, [float(x) if hasattr(x, "quantize") else x for x in r]))
        d["mcap_b"] = round(d["market_cap"] / 1e9, 1) if d.get("market_cap") else None
        rows.append(d)
    return rows


def add_sheet(wb, title, rows, extra=None):
    ws = wb.create_sheet(title)
    cols = COLS + (extra or [])
    for j, (hdr, _, w) in enumerate(cols, 1):
        cell = ws.cell(1, j, hdr)
        cell.font = Font(name="Calibri", bold=True, color="FFFFFF", size=10)
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(j)].width = w
    red = Font(name="Calibri", size=10, color="B54334", bold=True)
    for i, d in enumerate(rows, 2):
        for j, (_, key, _w) in enumerate(cols, 1):
            v = d.get(key)
            if isinstance(v, float):
                v = round(v, 4)
            cell = ws.cell(i, j, v)
            cell.font = Font(name="Calibri", size=10)
            if key == "comp" and v is not None and v < 6.5:
                cell.font = red
            if key == "earnings_quality_flag" and v == "EPS_ENGINEERED":
                cell.font = red
            if key in ("p2", "p3") and v is not None and v < 6.0:
                cell.font = red
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{len(rows)+1}"
    return ws


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", default=str(ROOT / "outputs" / "exports" /
                                         f"universe_and_portfolio_{date.today()}.xlsx"))
    a = ap.parse_args()
    conn = psycopg2.connect(settings.DATABASE_URL)
    cur = conn.cursor()
    rows = fetch(cur)

    holds = {}
    cur.execute("SELECT ticker, MAX(shares), MAX(avg_cost) FROM ic_portfolio_holdings GROUP BY ticker")
    for t, sh, ac in cur.fetchall():
        holds[t] = (float(sh) if sh else None, float(ac) if ac else None)
    conn.close()

    wb = Workbook()
    wb.remove(wb.active)
    add_sheet(wb, "Universe", rows)
    for q in ("Q1", "Q2", "Q3", "Q4"):
        qr = sorted([d for d in rows if d["quadrant"] == q],
                    key=lambda d: -(d["alignment_score_v3"] or 0))
        add_sheet(wb, q, qr)
    prows = []
    for d in rows:
        if not d["in_portfolio"]:
            continue
        sh, ac = holds.get(d["ticker"], (None, None))
        d = dict(d)
        d["shares"], d["avg_cost"] = sh, ac
        d["pl_pct"] = (round((d["current_price"] / ac - 1) * 100, 1)
                       if ac and d.get("current_price") else None)
        prows.append(d)
    prows.sort(key=lambda d: d["ticker"])
    add_sheet(wb, "Portfolio", prows,
              extra=[("Shares", "shares", 9), ("Avg Cost", "avg_cost", 9), ("P&L %", "pl_pct", 8)])
    wb.save(a.out)
    print(f"  {a.out} · {len(rows)} names · {len(prows)} holdings")


if __name__ == "__main__":
    main()
